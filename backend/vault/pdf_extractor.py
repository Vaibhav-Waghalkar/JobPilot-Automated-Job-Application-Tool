import io
from pathlib import Path
from pypdf import PdfReader
from docx import Document

def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        pages_text = []
        for p in reader.pages:
            txt = p.extract_text()
            if txt:
                pages_text.append(txt)
        return "\n".join(pages_text).strip()
    except Exception as e:
        return f"[Failed to extract PDF text: {str(e)}]"

def extract_text_from_docx(docx_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(docx_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs).strip()
    except Exception as e:
        return f"[Failed to extract DOCX text: {str(e)}]"

def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif suffix in [".docx", ".doc"]:
        return extract_text_from_docx(file_bytes)
    else:
        try:
            return file_bytes.decode("utf-8", errors="replace")
        except Exception:
            return ""
